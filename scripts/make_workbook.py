#!/usr/bin/env python3
"""Write and validate the research workbook using an existing Python installation.

The numerical analysis remains in the pipeline. This script reports typed result
values and checks every exported table cell against its original CSV field.
"""
from __future__ import annotations
import argparse
from collections import OrderedDict
from datetime import datetime,timezone
import hashlib,json,math,re,textwrap
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
from workbook_catalog import SECTIONS,LABELS,VALUE_LABELS

FONT='Arial'
INK='243746'
HEADER='284B63'
LIGHT='D8E8F1'
def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()
def serial(value):
    if value is None or (not isinstance(value,(str,bool)) and pd.isna(value)):return None
    if isinstance(value,np.generic):value=value.item()
    if isinstance(value,float):
        if not math.isfinite(value):return None
    return value
def label(field):return LABELS.get(field,field.replace('_',' ').capitalize())
def display(field,value):
    value=serial(value)
    if isinstance(value,bool):return 'Yes' if value else 'No'
    return VALUE_LABELS.get(field,{}).get(value,value) if isinstance(value,str) else value
def number_format(field,value):
    if isinstance(value,bool) or not isinstance(value,(int,float)):return 'General'
    if isinstance(value,int) or field in ('n','genes','n_genes','n_features','draws','n_total','n_finite','n_unavailable','n_target_pools','n_common_features','removed_count','remaining_count','n_at_least_observed'):
        return '#,##0'
    if ('fraction' in field or field.startswith('fraction_') or field in ('any_GO_coverage_of_original_panel',)):
        return '0.0%;-0.0%;0.0%'
    if value and abs(value)<1e-5:return '0.000E+00'
    return '0.000000;-0.000000;0.000000'
def table_width(field,values):
    longest=max([len(str(v)) for v in values if v is not None] or [0])
    if field in ('definition','rights','operation'):return min(115,max(75,longest*.35))
    if field in ('record_url','input_file'):return min(105,max(48,longest*.9))
    if field=='sha256':return 72
    if field=='topic':return 30
    if field in ('model','comparator','metric','category_name','score','latent_null','pairing_role','component'):
        return min(45,max(27,longest*.85))
    if field in ('source_file','source_rows','source_fields','operation','Source','Detail','sha256','path'):
        return min(74,max(32,longest*.72))
    if field in ('target_library','series'):return min(43,max(24,longest*.9))
    if field=='genotype':return 26
    if 'unavailable_reason' in field or field=='critical_prefix_status':return 35
    return min(29,max(17,len(label(field))*.48))

def write_table(ws,start,definition,frame,record_id,source_path):
    fields=definition['fields'];rows=[[display(f,row[f]) for f in fields] for _,row in frame.iterrows()]
    ncols=len(fields);endcol=get_column_letter(ncols)
    widths=[]
    for j,f in enumerate(fields,1):
        letter=get_column_letter(j)
        width=max(ws.column_dimensions[letter].width or 0,table_width(f,[row[j-1] for row in rows]))
        ws.column_dimensions[letter].width=width; widths.append(width)
    ws.cell(start,1,definition['title']).font=Font(name=FONT,size=14,bold=True,color=INK)
    note_lines=textwrap.wrap(definition['note'],width=155)
    for i,line in enumerate(note_lines):
        cell=ws.cell(start+2+i,1,line);cell.font=Font(name=FONT,size=10,italic=True,color='526371')
        cell.alignment=Alignment(vertical='center',wrap_text=False)
        ws.row_dimensions[start+2+i].height=17
    header_row=start+3+len(note_lines)
    for j,f in enumerate(fields,1):
        cell=ws.cell(header_row,j,label(f));cell.font=Font(name=FONT,size=10,bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor=HEADER)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        cell.border=Border(bottom=Side(style='thin',color='FFFFFF'))
    ws.row_dimensions[header_row].height=48
    for i,row in enumerate(rows,header_row+1):
        longest_lines=1
        for j,(f,value) in enumerate(zip(fields,row),1):
            cell=ws.cell(i,j,value)
            cell.font=Font(name=FONT,size=10,color=INK)
            wrap=isinstance(value,str) and (len(value)>max(12,widths[j-1]*.85) or '\n' in value)
            horizontal='right' if isinstance(value,(int,float)) else ('center' if value in ('Yes','No') else 'left')
            cell.alignment=Alignment(horizontal=horizontal,vertical='center',wrap_text=wrap,indent=0 if horizontal=='center' else 1)
            cell.number_format=number_format(f,value)
            if wrap:longest_lines=max(longest_lines,sum(max(1,math.ceil(len(v)/max(12,widths[j-1]*.9))) for v in value.split('\n')))
        ws.row_dimensions[i].height=min(210,max(21,16*longest_lines+4))
    last=header_row+len(rows)
    if rows:
        table=Table(displayName='ResultsTable'+str(record_id),ref=f'A{header_row}:{endcol}{last}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
        ws.add_table(table)
    # Source records are ordinary text beside/below the corresponding data.
    source_row=last+2
    ws.cell(source_row,1,'Source: '+source_path).font=Font(name=FONT,size=9,color='526371')
    ws.row_dimensions[source_row].height=16
    ws.print_options.horizontalCentered=False
    ws.print_options.verticalCentered=False
    return {'sheet':ws.title,'section':definition['title'],'source':source_path,'source_sha256':sha(frame.attrs['source_path']),
        'source_row_count':len(frame),'fields':fields,'header_row':header_row,'first_data_row':header_row+1,
        'last_data_row':last,'first_row':start,'last_row':source_row,'columns':ncols,'table_id':record_id},source_row+4

def figure_specs(directory):
    result=[]
    if directory is None:return result
    for i in range(1,5):
        p=directory/f'Figure_{i}_source.csv'
        if not p.is_file():raise FileNotFoundError('Missing main figure source: '+str(p))
        frame=pd.read_csv(p)
        preferred=['figure','panel','series','metric','genotype','comparator','value_axis','value','lower','upper','n','n_unit','n_features','range_type','span','denominator','x','y','source_file','source_rows','source_fields','operation']
        fields=[f for f in preferred if f in frame]+[f for f in frame if f not in preferred]
        if not fields:raise ValueError('No recognized figure-source columns: '+str(p))
        result.append({'sheet':f'Figure_{i}_data','title':f'Figure {i} numerical source',
            'path':str(p),'fields':fields,
            'note':'Exact plot records. Operation and source fields identify derived values. Ranges retain the definitions in the figure legend; no range is automatically a confidence interval.'})
    return result

def provenance_specs(output_directory):
    repo=Path(__file__).resolve().parents[1]
    source_manifest=repo/'config/data_sources.json'
    records=json.loads(source_manifest.read_text())['sources']
    rows=[]
    for rec in records:
        rows.append({'accession':rec['accession'],'input_file':Path(rec['relative_path']).name,
            'processed_layer':rec['kind'],'record_url':rec['record_url'],'sha256':rec['sha256'],
            'bytes':rec['bytes'],'rights':rec['rights']})
    source_csv=output_directory/'workbook_sources.csv'
    pd.DataFrame(rows).to_csv(source_csv,index=False)
    methods=[
     ('Study scope','Three GEO studies spanning Arabidopsis and citrus. Carrizo is a hybrid; the collection is described as two plant systems, not two species.'),
     ('Independent units','Three RNA pools per genotype/condition. Individual plants within a pool, genes, and overlapping resampling partitions are not additional independent RNA replicates.'),
     ('Discovery chronology','Four-factor donor selection followed earlier discovery outcomes. Numeric four-way targets were withheld in the documented prediction, but published qualitative behavior was known.'),
     ('External-study scope','Prediction rules were fixed before numeric target access. Investigators overlap across studies, so this is not fully independent-laboratory replication.'),
     ('Preprocessing','Reanalysis starts from deposited processed expression matrices. It does not rerun raw FASTQ alignment or remove information coupling introduced by original preprocessing.'),
     ('Primary settings','Primary external Arabidopsis endpoint: wild type at mean FPKM filter 1. Both citrus genotypes use filter 1. Other filters and the Arabidopsis mutant are sensitivity/secondary settings.'),
     ('Prediction rule','The same arithmetic rule is fitted separately to each genotype’s lower-order observations. No fitted model is transferred across genotypes or species.'),
     ('Observed response spans','Projection residuals concern estimated noisy spans. They are not pure higher-order interactions and do not establish a new biological response dimension.'),
     ('Conditional simulations','All 36 conditional scenarios are retained. Nuisance mean/covariance choices use known outcomes. Exceedance fractions are not calibrated biological P values; the true biological span remains unidentified.'),
     ('Gene gains','Per-gene gains are comparator-relative differences of prediction losses. They are not causal gene effects or differential-expression statistics.'),
     ('Functional groups','Exact-Entrez GO annotations use computational IEA evidence. Groups overlap and some are small. No universal pathway, activation or significant enrichment is claimed.'),
     ('Complementary refits','256 preselected overlapping partitions per genotype. The two sides in one partition are disjoint; the 1,024 total fits are not mutually independent.'),
     ('Uncertainty labels','Refit quantiles are partition-sensitivity ranges, not confidence intervals. No gene bootstrap or gene-as-biological-replicate P value is used.'),
     ('Local validation scope','Portable Python numerical and file validation is recorded separately. Successful export does not establish native Microsoft Excel or Windows PowerShell execution.'),
    ]
    methods_csv=output_directory/'workbook_methods.csv'
    pd.DataFrame(methods,columns=['topic','definition']).to_csv(methods_csv,index=False)
    return [dict(sheet='Sources',title='Processed input provenance',path=str(source_csv),
        fields=['accession','input_file','processed_layer','record_url','bytes','sha256','rights'],
        note='Original deposited data retain their source rights and attribution. Code licensing does not relicense third-party data.'),
        dict(sheet='Sources',title='Study definitions and interpretation limits',path=str(methods_csv),
        fields=['topic','definition'],note='These definitions apply to the displayed result tables and their figure counterparts.')]

def validate_export(workbook_path,records,results_root):
    wb=load_workbook(workbook_path,read_only=False,data_only=False)
    counts={'numeric':0,'text':0,'boolean':0,'blank':0};sources={};sections=[]
    for rec in records:
        path=Path(rec['source_absolute'])
        if sha(path)!=rec['source_sha256']:raise ValueError('Source changed during workbook export: '+str(path))
        frame=pd.read_csv(path)
        if len(frame)!=rec['source_row_count']:raise ValueError('Row count mismatch')
        ws=wb[rec['sheet']]
        for i,(_,row) in enumerate(frame.iterrows(),rec['first_data_row']):
            for j,field in enumerate(rec['fields'],1):
                expected=display(field,row[field]);actual=ws.cell(i,j).value
                key=f"{rec['sheet']}!{get_column_letter(j)}{i}"
                if expected is None:
                    assert actual is None,key;counts['blank']+=1
                elif isinstance(expected,bool):
                    assert isinstance(actual,bool) and actual==expected,key;counts['boolean']+=1
                elif isinstance(expected,(int,float)):
                    assert isinstance(actual,(int,float)) and not isinstance(actual,bool),key
                    assert math.isclose(actual,expected,rel_tol=2e-12,abs_tol=2e-12),(key,actual,expected)
                    counts['numeric']+=1
                else:
                    assert actual==expected,(key,actual,expected);counts['text']+=1
        sources[rec['source']]=rec['source_sha256']
        sections.append({k:rec[k] for k in ('sheet','section','source','source_row_count','first_data_row','last_data_row','columns')})
    errors=[];formulas=[]
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type=='e':errors.append(f'{ws.title}!{cell.coordinate}')
                if cell.data_type=='f':formulas.append(f'{ws.title}!{cell.coordinate}')
    assert not errors,errors
    assert not formulas,'Scientific outputs must be reported from executed CSVs, not reimplemented as Excel formulas.'
    wb.close()
    return {'status':'PASS','workbook_sha256':sha(workbook_path),'workbook_bytes':workbook_path.stat().st_size,
        'cells_verified_directly_against_csv':sum(counts.values()),'verified_cells_by_type':counts,
        'excel_error_cells':len(errors),'formula_cells':len(formulas),'sources_unchanged':True,
        'source_hashes':sources,'sections':sections,
        'label_mappings':'Scientific identifiers use the published label dictionary; source Boolean values display as Yes/No. Numeric values are not transformed by this exporter.',
        'validation_scope':'Fresh exported workbook values and labels checked directly against all displayed CSV source fields. Rendering and native Excel execution are separate checks.'}

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--results',type=Path,required=True,help='Fresh analysis output containing tables/.')
    parser.add_argument('--output',type=Path,required=True,help='New workbook path.')
    parser.add_argument('--figure-sources',type=Path,help='Directory containing the four main Figure_N_source.csv files.')
    args=parser.parse_args();root=args.results.resolve();out=args.output.resolve()
    if out.exists():raise FileExistsError('Use a fresh workbook output path: '+str(out))
    out.parent.mkdir(parents=True,exist_ok=True)
    definitions=[dict(x) for x in SECTIONS]+figure_specs(args.figure_sources.resolve() if args.figure_sources else None)+provenance_specs(out.parent)
    wb=Workbook();wb.remove(wb.active);records=[];positions={}
    for i,definition in enumerate(definitions,1):
        path=Path(definition['path'])
        if not path.is_absolute():path=root/'tables'/path
        if not path.is_file():raise FileNotFoundError('Required analysis table is missing: '+str(path))
        frame=pd.read_csv(path)
        missing=set(definition['fields'])-set(frame)
        if missing:raise ValueError(f'{path.name}: missing fields {sorted(missing)}')
        frame.attrs['source_path']=path
        name=definition['sheet']
        if name not in wb.sheetnames:
            ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
            ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.sheet_properties.outlinePr.summaryRight=True
            ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
            ws.page_setup.fitToHeight=0;ws.page_setup.fitToWidth=1
            ws.sheet_properties.tabColor=HEADER if len(wb.sheetnames)<=3 else None
            positions[name]=2
        ws=wb[name]
        try:source=str(path.relative_to(root))
        except ValueError:source=('figure_source/' if path.name.startswith('Figure_') else 'report_metadata/')+path.name
        rec,positions[name]=write_table(ws,positions[name],definition,frame,i,source)
        rec['source_absolute']=str(path);records.append(rec)
    for ws in wb:
        own=[r for r in records if r['sheet']==ws.title]
        first=own[0]
        if ws.max_row>24:ws.freeze_panes=f"B{first['header_row']+1}"
        ws.print_area=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.print_title_rows=f"1:{first['header_row']}"
    wb.properties.title='Stress-combination prediction results'
    wb.properties.subject='Executed predictions, uncertainty diagnostics, and gene-gain repeatability'
    wb.properties.creator='Research analysis pipeline'
    wb.properties.description='Three GEO studies spanning Arabidopsis and citrus. Scientific values are traced to executed CSV tables.'
    wb.save(out);wb.close()
    validation=validate_export(out,records,root)
    validation['created_utc']=datetime.now(timezone.utc).isoformat()
    validation['sheets']=list(dict.fromkeys(r['sheet'] for r in records))
    validation['main_figure_source_sheets_included']=args.figure_sources is not None
    validation_path=out.parent/'workbook_validation.json'
    validation_path.write_text(json.dumps(validation,indent=2,allow_nan=False)+'\n')
    trace=[{k:v for k,v in rec.items() if k!='source_absolute'} for rec in records]
    (out.parent/'workbook_cell_lineage.json').write_text(json.dumps(trace,indent=2)+'\n')
    print(json.dumps({'status':validation['status'],'workbook':str(out),'sheets':len(validation['sheets']),
        'cells_verified_directly_against_csv':validation['cells_verified_directly_against_csv'],
        'validation':str(validation_path)},indent=2))

if __name__=='__main__':main()
